"""Build the NEMT risk-sweep trade workbook (Excel) from trades_all.json."""

from __future__ import annotations

import json
import statistics
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SOURCE = Path(r"C:\Users\Administrator\AppData\Local\Temp\trades_all.json")
TARGET = Path(r"E:\NEMT Simulator\docs\reports\NEMT_trades_risk_sweep.xlsx")

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
TITLE_FONT = Font(bold=True, size=12, color="1F3864")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MONEY = "#,##0.00"
PRICE = "#,##0.0000"
PCT = "0.00"
THOUS = "#,##0.00"


def style_header(ws, row: int, columns: int) -> None:
    for col in range(1, columns + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = 30


def write_table(ws, headers: list[str], rows: list[dict[str, Any]],
                formats: dict[str, str] | None = None,
                widths: dict[str, int] | None = None) -> None:
    formats = formats or {}
    for index, name in enumerate(headers, 1):
        ws.cell(row=1, column=index, value=name)
    style_header(ws, 1, len(headers))
    for r, record in enumerate(rows, 2):
        for c, name in enumerate(headers, 1):
            cell = ws.cell(row=r, column=c, value=record.get(name))
            if name in formats and isinstance(cell.value, (int, float)):
                cell.number_format = formats[name]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    for index, name in enumerate(headers, 1):
        width = (widths or {}).get(name, max(10, min(22, len(name) + 4)))
        ws.column_dimensions[get_column_letter(index)].width = width


def streaks(pnls: list[float]) -> tuple[int, float]:
    worst = cur = 0
    for value in pnls:
        if value <= 0:
            cur += 1
            worst = max(worst, cur)
        else:
            cur = 0
    return worst, worst


def main() -> None:
    payload = json.loads(SOURCE.read_text(encoding="utf-8"))
    runs: list[dict[str, Any]] = payload["runs"]
    trades: list[dict[str, Any]] = payload["trades"]

    wb = Workbook()

    # ---------- 1. README ----------
    ws = wb.active
    ws.title = "说明"
    ws["A1"] = "NEMT 策略 · 风险预算敏感性回测 · 全量开单记录"
    ws["A1"].font = Font(bold=True, size=14, color="1F3864")
    notes = [
        ("文件用途", "把 64 次回测 × 6 个策略格子的每一笔开单全部导出，便于做统计与透视"),
        ("数据来源", "Binance 现货 4h 官方归档（校验 .CHECKSUM）；BTCUSDT 与 ETHUSDT"),
        ("区间", "2022-01-01 至 2026-01-01，按自然年独立运行（每年年初重置 100 USDT）"),
        ("风险档", "单笔风险 1% / 2% / 3% / 5% / 7.5% / 10%"),
        ("杠杆", "3x 为主；另有 5% 与 10% 风险在 10x 下的算术演示档（无强平模型，不可交易）"),
        ("格子", "R0 单笔 / F0 固定预算 / T3 三批 / T2 两批 / C1 平衡止损 / U1 周级UP-UP闸门"),
        ("交易总数", f"{len(trades)} 笔；配置总数 {len(runs)} 个"),
        ("R 倍数", "net_pnl ÷ (风险比例 × 开仓前权益)，即每笔以初始风险为单位的盈亏"),
        ("证据等级", "EXPLORATORY —— 所有年份与资产均已被多轮研究使用过，不可作确认性结论"),
        ("重要限制", "未计资金费率、强平、延迟、盘口深度；10x 档在真实交易中会被强平"),
        ("订单权限", "无。本工作簿不涉及任何 PAPER / TESTNET / LIVE 订单"),
    ]
    for r, (key, value) in enumerate(notes, 3):
        ws.cell(row=r, column=1, value=key).font = Font(bold=True)
        cell = ws.cell(row=r, column=2, value=value)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 30
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 110

    # ---------- 2. 总览 (run x cell) ----------
    ws = wb.create_sheet("总览")
    run_headers = ["run_id", "symbol", "year", "risk_pct", "max_leverage", "cell",
                   "final", "ret_pct", "dd_pct", "trades", "wins", "losses",
                   "profit_factor", "max_open_risk_pct", "leverage_truncations",
                   "w_o_best_pct", "fees"]
    run_formats = {"final": MONEY, "ret_pct": PCT, "dd_pct": PCT, "profit_factor": "0.000",
                   "max_open_risk_pct": PCT, "w_o_best_pct": PCT, "fees": "0.0000",
                   "risk_pct": "0.0"}
    run_rows = sorted(runs, key=lambda x: (x["symbol"], x["year"], x["risk_pct"],
                                           x["max_leverage"], x["cell"]))
    write_table(ws, run_headers, run_rows, run_formats,
                {"run_id": 22, "symbol": 11, "cell": 8, "profit_factor": 12})

    # ---------- 3. 全部交易 ----------
    ws = wb.create_sheet("全部交易")
    trade_headers = ["run_id", "symbol", "year", "risk_pct", "max_leverage", "cell",
                     "seq", "side", "exit_reason", "exit_at", "holding_bars",
                     "holding_days", "avg_entry", "exit_price", "quantity", "notional",
                     "entry_fee", "exit_fee", "net_pnl", "equity_before", "equity_after",
                     "risk_usdt", "r_multiple", "tranches"]
    trade_formats = {"avg_entry": PRICE, "exit_price": PRICE, "notional": THOUS,
                     "entry_fee": "0.0000", "exit_fee": "0.0000", "net_pnl": "0.0000",
                     "equity_before": MONEY, "equity_after": MONEY, "risk_usdt": "0.0000",
                     "r_multiple": "0.000", "holding_days": "0.00", "risk_pct": "0.0"}
    trade_rows = sorted(trades, key=lambda x: (x["symbol"], x["year"], x["risk_pct"],
                                               x["max_leverage"], x["cell"], x["seq"]))
    write_table(ws, trade_headers, trade_rows, trade_formats,
                {"run_id": 22, "symbol": 11, "exit_at": 20, "exit_reason": 15,
                 "side": 7, "cell": 8})

    # ---------- 4. 风险档汇总 ----------
    ws = wb.create_sheet("风险档汇总")
    buckets: dict[tuple, list[dict[str, Any]]] = defaultdict(list)
    for row in trades:
        buckets[(row["symbol"], row["max_leverage"], row["risk_pct"], row["cell"])].append(row)
    summary_rows = []
    for key in sorted(buckets, key=lambda k: (k[0], k[1], k[2], k[3])):
        sym, lev, risk, cell = key
        items = buckets[key]
        rs = [x["r_multiple"] for x in items if x["r_multiple"] is not None]
        wins = [x["net_pnl"] for x in items if x["net_pnl"] > 0]
        losses = [-x["net_pnl"] for x in items if x["net_pnl"] <= 0]
        gross_win, gross_loss = sum(wins), sum(losses)
        summary_rows.append({
            "symbol": sym, "max_leverage": lev, "risk_pct": risk, "cell": cell,
            "trades": len(items),
            "win_rate_pct": round(len(wins) / len(items) * 100, 2) if items else 0,
            "avg_R": round(statistics.fmean(rs), 3) if rs else None,
            "median_R": round(statistics.median(rs), 3) if rs else None,
            "best_R": round(max(rs), 3) if rs else None,
            "worst_R": round(min(rs), 3) if rs else None,
            "expectancy_R": round(sum(rs) / len(rs), 3) if rs else None,
            "sum_R": round(sum(rs), 2) if rs else None,
            "profit_factor": round(gross_win / gross_loss, 3) if gross_loss else None,
            "avg_net_pnl": round(statistics.fmean(x["net_pnl"] for x in items), 4),
            "avg_holding_days": round(statistics.fmean(x["holding_days"] for x in items), 2),
            "fees_total": round(sum(x["entry_fee"] + x["exit_fee"] for x in items), 4),
        })
    summary_headers = ["symbol", "max_leverage", "risk_pct", "cell", "trades",
                       "win_rate_pct", "avg_R", "median_R", "best_R", "worst_R",
                       "expectancy_R", "sum_R", "profit_factor", "avg_net_pnl",
                       "avg_holding_days", "fees_total"]
    write_table(ws, summary_headers, summary_rows,
                {"win_rate_pct": PCT, "avg_R": "0.000", "median_R": "0.000",
                 "best_R": "0.000", "worst_R": "0.000", "expectancy_R": "0.000",
                 "sum_R": "0.00", "profit_factor": "0.000", "avg_net_pnl": "0.0000",
                 "avg_holding_days": "0.00", "fees_total": "0.0000", "risk_pct": "0.0"})

    # ---------- 5. 年份汇总 (all cells pooled per symbol/year/risk) ----------
    ws = wb.create_sheet("年份汇总")
    year_buckets: dict[tuple, list[dict[str, Any]]] = defaultdict(list)
    for row in trades:
        year_buckets[(row["symbol"], row["year"], row["max_leverage"],
                      row["risk_pct"])].append(row)
    year_rows = []
    for key in sorted(year_buckets, key=lambda k: (k[0], k[1], k[2], k[3])):
        sym, year, lev, risk = key
        items = year_buckets[key]
        rs = [x["r_multiple"] for x in items if x["r_multiple"] is not None]
        losses = [x for x in items if x["net_pnl"] <= 0]
        streak, _ = streaks([x["net_pnl"] for x in items])
        year_rows.append({
            "symbol": sym, "year": year, "max_leverage": lev, "risk_pct": risk,
            "cells_pooled": len({x["cell"] for x in items}),
            "trades": len(items),
            "win_rate_pct": round((len(items) - len(losses)) / len(items) * 100, 2),
            "avg_R": round(statistics.fmean(rs), 3) if rs else None,
            "sum_R": round(sum(rs), 2) if rs else None,
            "max_consec_losses": streak,
            "avg_holding_days": round(statistics.fmean(x["holding_days"] for x in items), 2),
        })
    write_table(ws, ["symbol", "year", "max_leverage", "risk_pct", "cells_pooled",
                     "trades", "win_rate_pct", "avg_R", "sum_R",
                     "max_consec_losses", "avg_holding_days"], year_rows,
                {"win_rate_pct": PCT, "avg_R": "0.000", "sum_R": "0.00",
                 "avg_holding_days": "0.00", "risk_pct": "0.0"})

    # ---------- 6. 连亏分析 (per run x cell) ----------
    ws = wb.create_sheet("连亏分析")
    run_trades: dict[tuple, list[dict[str, Any]]] = defaultdict(list)
    for row in trades:
        run_trades[(row["run_id"], row["cell"])].append(row)
    streak_rows = []
    for key in sorted(run_trades):
        run_id, cell = key
        items = sorted(run_trades[key], key=lambda x: x["seq"])
        pnls = [x["net_pnl"] for x in items]
        worst_count, _ = streaks(pnls)
        worst_loss = min(pnls) if pnls else 0.0
        streak_rows.append({
            "run_id": run_id, "cell": cell, "trades": len(items),
            "wins": sum(1 for p in pnls if p > 0),
            "losses": sum(1 for p in pnls if p <= 0),
            "max_consec_losses": worst_count,
            "worst_single_trade": round(worst_loss, 4),
            "worst_single_R": round(min(x["r_multiple"] for x in items
                                        if x["r_multiple"] is not None), 3),
            "best_single_trade": round(max(pnls), 4),
            "best_single_R": round(max(x["r_multiple"] for x in items
                                       if x["r_multiple"] is not None), 3),
            "net_total": round(sum(pnls), 4),
        })
    write_table(ws, ["run_id", "cell", "trades", "wins", "losses",
                     "max_consec_losses", "worst_single_trade", "worst_single_R",
                     "best_single_trade", "best_single_R", "net_total"], streak_rows,
                {"worst_single_trade": "0.0000", "worst_single_R": "0.000",
                 "best_single_trade": "0.0000", "best_single_R": "0.000",
                 "net_total": "0.0000"}, {"run_id": 22, "cell": 8})

    # ---------- 7. 退出原因 ----------
    ws = wb.create_sheet("退出原因")
    reasons: dict[tuple, Counter] = defaultdict(Counter)
    for row in trades:
        reasons[(row["symbol"], row["exit_reason"])][row["cell"]] += 1
    reason_rows = []
    for (sym, reason), counter in sorted(reasons.items()):
        record = {"symbol": sym, "exit_reason": reason, "total": sum(counter.values())}
        record.update({f"cell_{k}": v for k, v in sorted(counter.items())})
        reason_rows.append(record)
    cells_present = sorted({k for counter in reasons.values() for k in counter})
    reason_headers = ["symbol", "exit_reason", "total"] + [f"cell_{c}" for c in cells_present]
    write_table(ws, reason_headers, reason_rows, None, {"exit_reason": 18})

    TARGET.parent.mkdir(parents=True, exist_ok=True)
    wb.save(TARGET)
    print(f"saved {TARGET} ({TARGET.stat().st_size:,} bytes)")
    print(f"sheets: {wb.sheetnames}")
    print(f"trades={len(trades)} runs={len(runs)}")


if __name__ == "__main__":
    main()
